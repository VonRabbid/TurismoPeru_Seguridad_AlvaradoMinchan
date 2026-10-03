"""
================================================================================
PROYECTO: TURISMOPERÚ - MÓDULO ANALÍTICO DE INTELIGENCIA DE NEGOCIOS
EVALUACIÓN: Tercera Evaluación de Base de Datos II - Grupo B (Alternativa B)
INSTITUCIÓN: Universidad Nacional de Cajamarca (UNC)
AUTOR: Cristian Paul Alvarado Minchan
BASE DE DATOS: TURISMOPERU_CPAM
================================================================================
"""

import os
import sys
from datetime import datetime
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import seaborn as sns
from dotenv import load_dotenv
import sqlalchemy as sa
from sqlalchemy import create_engine
import pyodbc

# ReportLab imports para compilación del informe ejecutivo en PDF
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# python-docx imports para compilación del informe ejecutivo en DOCX
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# ------------------------------------------------------------------------------
# 1. CONFIGURACIÓN DE ENTORNO Y RUTAS
# ------------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))

# Cargar variables de entorno buscando en 06_python/.env o en la raíz
env_path = os.path.join(SCRIPT_DIR, ".env")
if not os.path.exists(env_path):
    env_path = os.path.join(PROJECT_ROOT, ".env")
load_dotenv(dotenv_path=env_path)

DB_SERVER = os.getenv("DB_SERVER", "161.132.54.162")
DB_DATABASE = os.getenv("DB_DATABASE", "TURISMOPERU_CPAM")
DB_USER = os.getenv("DB_USER", "")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_DRIVER = os.getenv("DB_DRIVER", "ODBC Driver 18 for SQL Server")

EVIDENCIAS_DIR = os.path.join(PROJECT_ROOT, "evidencias")
REPORTES_DIR = os.path.join(PROJECT_ROOT, "05_reportes")
TEMP_IMG_DIR = os.path.join(SCRIPT_DIR, "temp_charts")

os.makedirs(EVIDENCIAS_DIR, exist_ok=True)
os.makedirs(REPORTES_DIR, exist_ok=True)
os.makedirs(TEMP_IMG_DIR, exist_ok=True)

# ------------------------------------------------------------------------------
# 2. CONEXIÓN A BASE DE DATOS MEDIANTE SQLALCHEMY / PYODBC
# ------------------------------------------------------------------------------
def get_db_engine():
    if not DB_USER or not DB_PASSWORD:
        raise ValueError(
            "Credenciales de base de datos no configuradas. "
            "Por favor, configure DB_USER y DB_PASSWORD en el archivo .env"
        )
    print(f"[CONEXIÓN] Conectando a {DB_SERVER} -> Base de datos: {DB_DATABASE}...")
    connection_url = sa.engine.URL.create(
        "mssql+pyodbc",
        username=DB_USER,
        password=DB_PASSWORD,
        host=DB_SERVER,
        database=DB_DATABASE,
        query={
            "driver": DB_DRIVER,
            "Encrypt": "yes",
            "TrustServerCertificate": "yes"
        }
    )
    engine = create_engine(connection_url, fast_executemany=True)
    return engine

# ------------------------------------------------------------------------------
# 3. EXTRACCIÓN Y TRANSFORMACIÓN DE DATOS (ETL)
# ------------------------------------------------------------------------------
def extract_data(engine):
    print("[ETL] Consultando tablas transaccionales de TURISMOPERU_CPAM...")
    
    # Clientes con datos personales
    query_clientes = """
    SELECT 
        c.id_persona,
        p.numero_documento,
        p.nombres,
        p.apaterno,
        p.amaterno,
        RTRIM(p.nombres) + ' ' + RTRIM(p.apaterno) AS cliente_nombre,
        p.estado,
        p.fecha_registro
    FROM CPAM.cliente c
    JOIN CPAM.persona p ON c.id_persona = p.id_persona;
    """
    df_clientes = pd.read_sql(query_clientes, engine)

    # Reservas con sus estados y precios
    query_reservas = """
    SELECT 
        r.id_reserva,
        r.codigo_reserva,
        r.id_cliente,
        r.precio_total,
        r.adelanto,
        r.saldo_pendiente,
        r.fecha_reserva,
        r.fecha_inicio,
        r.fecha_fin,
        r.numero_personas,
        ISNULL(er.nombre, 'Sin Estado') AS estado_reserva
    FROM CPAM.reserva r
    LEFT JOIN CPAM.estado_reserva er ON r.id_estado_reserva = er.id_estado_reserva;
    """
    df_reservas = pd.read_sql(query_reservas, engine)
    df_reservas['fecha_reserva'] = pd.to_datetime(df_reservas['fecha_reserva'])
    df_reservas['mes_periodo'] = df_reservas['fecha_reserva'].dt.strftime('%Y-%m')

    # Pagos con sus medios de pago
    query_pagos = """
    SELECT 
        p.id_pago,
        p.id_reserva,
        p.id_medio_pago,
        p.monto,
        p.fecha_pago,
        p.estado AS estado_pago,
        ISNULL(mp.nombre, 'Otro') AS medio_pago
    FROM CPAM.pago p
    LEFT JOIN CPAM.medio_pago mp ON p.id_medio_pago = mp.id_medio_pago;
    """
    df_pagos = pd.read_sql(query_pagos, engine)
    df_pagos['fecha_pago'] = pd.to_datetime(df_pagos['fecha_pago'])

    print(f" -> Registros extraídos: Clientes={len(df_clientes)}, Reservas={len(df_reservas)}, Pagos={len(df_pagos)}")
    return df_clientes, df_reservas, df_pagos

# ------------------------------------------------------------------------------
# 4. CÁLCULO DE KPIS
# ------------------------------------------------------------------------------
def compute_kpis(df_clientes, df_reservas, df_pagos):
    total_clientes = len(df_clientes)
    total_reservas = len(df_reservas)
    total_ingresos_pagos = float(df_pagos['monto'].sum())
    total_facturado_reservas = float(df_reservas['precio_total'].sum())
    ticket_promedio = total_ingresos_pagos / total_reservas if total_reservas > 0 else 0.0

    kpis = {
        "Total Clientes": total_clientes,
        "Total Reservas": total_reservas,
        "Total Ingresos": total_ingresos_pagos,
        "Ticket Promedio": ticket_promedio,
        "Facturación Contratada": total_facturado_reservas
    }

    print("\n=======================================================")
    print("                RESUMEN DE KPIS EJECUTIVOS             ")
    print("=======================================================")
    print(f"1. Total Clientes:               {total_clientes:,}")
    print(f"2. Total Reservas:               {total_reservas:,}")
    print(f"3. Total Ingresos Cobrados:      S/. {total_ingresos_pagos:,.2f}")
    print(f"4. Ticket Promedio (Cobrado):    S/. {ticket_promedio:,.2f}")
    print(f"   (Facturación Contratada):     S/. {total_facturado_reservas:,.2f}")
    print("=======================================================\n")
    return kpis

# ------------------------------------------------------------------------------
# 5. GENERACIÓN DE GRÁFICOS ANALÍTICOS
# ------------------------------------------------------------------------------
def generate_charts(df_clientes, df_reservas, df_pagos):
    print("[VISUALIZACIÓN] Generando gráficos estadísticos obligatorios...")
    plt.style.use('seaborn-v0_8-whitegrid')
    plt.rcParams['font.family'] = 'DejaVu Sans'
    
    palette_primary = "#1B365D"
    palette_secondary = "#008080"
    palette_accent = "#E65100"

    # --------------------------------------------------------------------------
    # Gráfico 1: Reservas por Estado
    # --------------------------------------------------------------------------
    fig1, ax1 = plt.subplots(figsize=(9, 5), dpi=300)
    estado_counts = df_reservas['estado_reserva'].value_counts()
    colors_est = sns.color_palette("mako", len(estado_counts))
    bars = ax1.barh(estado_counts.index[::-1], estado_counts.values[::-1], color=colors_est[::-1], edgecolor='black', linewidth=0.6)
    ax1.set_title("Distribución de Reservas por Estado Operativo", fontsize=14, fontweight='bold', pad=15, color='#1B365D')
    ax1.set_xlabel("Cantidad de Reservas", fontsize=11, fontweight='bold')
    ax1.set_ylabel("Estado de Reserva", fontsize=11, fontweight='bold')
    for bar in bars:
        w = bar.get_width()
        pct = (w / len(df_reservas)) * 100
        ax1.text(w + 0.6, bar.get_y() + bar.get_height()/2, f"{int(w)} ({pct:.1f}%)", 
                 va='center', fontsize=9, fontweight='bold', color='#2C3E50')
    ax1.set_xlim(0, max(estado_counts.values) + 5)
    plt.tight_layout()
    chart1_path = os.path.join(TEMP_IMG_DIR, "chart1_estados.png")
    fig1.savefig(chart1_path)
    plt.close(fig1)

    # --------------------------------------------------------------------------
    # Gráfico 2: Ingresos por Medio de Pago
    # --------------------------------------------------------------------------
    fig2, ax2 = plt.subplots(figsize=(9, 5), dpi=300)
    pagos_por_medio = df_pagos.groupby('medio_pago')['monto'].sum().sort_values(ascending=False)
    colors_mp = ["#1B365D", "#008080", "#2E7D32", "#F57C00", "#7B1FA2", "#C2185B"]
    bars2 = ax2.bar(pagos_por_medio.index, pagos_por_medio.values, color=colors_mp[:len(pagos_por_medio)], edgecolor='black', linewidth=0.6)
    ax2.set_title("Ingresos Recaudados por Medio de Pago (S/.)", fontsize=14, fontweight='bold', pad=15, color='#1B365D')
    ax2.set_ylabel("Monto Total Recaudado (Soles)", fontsize=11, fontweight='bold')
    ax2.set_xticks(range(len(pagos_por_medio.index)))
    ax2.set_xticklabels(pagos_por_medio.index, rotation=25, ha='right', fontsize=10, fontweight='bold')
    for bar in bars2:
        h = bar.get_height()
        pct = (h / pagos_por_medio.sum()) * 100
        ax2.text(bar.get_x() + bar.get_width()/2, h + 3000, f"S/. {h:,.0f}\n({pct:.1f}%)", 
                 ha='center', va='bottom', fontsize=9, fontweight='bold', color='#1B365D')
    ax2.set_ylim(0, max(pagos_por_medio.values) * 1.18)
    plt.tight_layout()
    chart2_path = os.path.join(TEMP_IMG_DIR, "chart2_medios_pago.png")
    fig2.savefig(chart2_path)
    plt.close(fig2)

    # --------------------------------------------------------------------------
    # Gráfico 3: Reservas por Período / Fecha (Tendencia Mensual)
    # --------------------------------------------------------------------------
    fig3, ax3 = plt.subplots(figsize=(9, 5), dpi=300)
    reservas_periodo = df_reservas.groupby('mes_periodo').size()
    ax3.plot(reservas_periodo.index, reservas_periodo.values, marker='o', markersize=8, 
             color='#1B365D', linewidth=2.5, label='Reservas por Mes')
    ax3.bar(reservas_periodo.index, reservas_periodo.values, color='#80DEEA', alpha=0.5, 
            edgecolor='#00838F', linewidth=1, width=0.4)
    ax3.set_title("Evolución Temporal de Reservas por Mes (2026)", fontsize=14, fontweight='bold', pad=15, color='#1B365D')
    ax3.set_xlabel("Período Mensual", fontsize=11, fontweight='bold')
    ax3.set_ylabel("Volumen de Reservas", fontsize=11, fontweight='bold')
    for idx, (p, val) in enumerate(reservas_periodo.items()):
        ax3.text(idx, val + 1.2, f"{val} res.", ha='center', fontsize=10, fontweight='bold', color='#004D40')
    ax3.set_ylim(0, max(reservas_periodo.values) + 5)
    plt.tight_layout()
    chart3_path = os.path.join(TEMP_IMG_DIR, "chart3_periodo.png")
    fig3.savefig(chart3_path)
    plt.close(fig3)

    # --------------------------------------------------------------------------
    # Gráfico 4: Top 10 Clientes con Mayor Cantidad de Reservas
    # --------------------------------------------------------------------------
    fig4, ax4 = plt.subplots(figsize=(9, 5), dpi=300)
    top_reservas_cli = df_reservas.groupby('id_cliente').size().reset_index(name='num_reservas')
    top_reservas_cli = top_reservas_cli.merge(df_clientes[['id_persona', 'cliente_nombre']], left_on='id_cliente', right_on='id_persona')
    top_reservas_cli = top_reservas_cli.sort_values(by=['num_reservas'], ascending=False).head(10)
    
    bars4 = ax4.barh(top_reservas_cli['cliente_nombre'][::-1], top_reservas_cli['num_reservas'][::-1], 
                     color='#2E7D32', edgecolor='black', linewidth=0.6)
    ax4.set_title("Top 10 Clientes por Frecuencia de Reservas", fontsize=14, fontweight='bold', pad=15, color='#1B365D')
    ax4.set_xlabel("Cantidad de Reservas Realizadas", fontsize=11, fontweight='bold')
    for bar in bars4:
        w = bar.get_width()
        ax4.text(w + 0.1, bar.get_y() + bar.get_height()/2, f"{int(w)} reservas", 
                 va='center', fontsize=9, fontweight='bold', color='#1B5E20')
    ax4.set_xlim(0, max(top_reservas_cli['num_reservas']) + 1)
    plt.tight_layout()
    chart4_path = os.path.join(TEMP_IMG_DIR, "chart4_top_clientes_reservas.png")
    fig4.savefig(chart4_path)
    plt.close(fig4)

    # --------------------------------------------------------------------------
    # Gráfico 5: Ingresos Totales Generados por Cliente (Top 10 por Facturación)
    # --------------------------------------------------------------------------
    fig5, ax5 = plt.subplots(figsize=(9, 5), dpi=300)
    ingresos_cli = df_reservas.groupby('id_cliente')['precio_total'].sum().reset_index(name='total_facturado')
    ingresos_cli = ingresos_cli.merge(df_clientes[['id_persona', 'cliente_nombre']], left_on='id_cliente', right_on='id_persona')
    ingresos_cli = ingresos_cli.sort_values(by='total_facturado', ascending=False).head(10)

    bars5 = ax5.barh(ingresos_cli['cliente_nombre'][::-1], ingresos_cli['total_facturado'][::-1], 
                     color='#E65100', edgecolor='black', linewidth=0.6)
    ax5.set_title("Top 10 Clientes con Mayor Valor de Facturación (S/.)", fontsize=14, fontweight='bold', pad=15, color='#1B365D')
    ax5.set_xlabel("Total Facturado (Soles)", fontsize=11, fontweight='bold')
    for bar in bars5:
        w = bar.get_width()
        ax5.text(w + 300, bar.get_y() + bar.get_height()/2, f"S/. {w:,.0f}", 
                 va='center', fontsize=9, fontweight='bold', color='#BF360C')
    ax5.set_xlim(0, max(ingresos_cli['total_facturado']) * 1.15)
    plt.tight_layout()
    chart5_path = os.path.join(TEMP_IMG_DIR, "chart5_ingresos_cliente.png")
    fig5.savefig(chart5_path)
    plt.close(fig5)

    # --------------------------------------------------------------------------
    # DASHBOARD CONSOLIDADO COMPLETO (evidencias/reporte.png)
    # --------------------------------------------------------------------------
    print("[DASHBOARD] Creando panel analítico integral en evidencias/reporte.png...")
    fig_all = plt.figure(figsize=(20, 14), dpi=200)
    fig_all.patch.set_facecolor('#F8F9FA')

    # Título superior
    plt.suptitle("TURISMOPERÚ - DASHBOARD EJECUTIVO DE GESTIÓN Y ANALÍTICA DE RESERVAS\n"
                 "Universidad Nacional de Cajamarca | Autor: Cristian Paul Alvarado Minchan | Base: TURISMOPERU_CPAM", 
                 fontsize=18, fontweight='bold', color='#1B365D', y=0.98)

    # Subplot 1: Estados
    ax_d1 = plt.subplot2grid((2, 3), (0, 0))
    ax_d1.barh(estado_counts.index[::-1], estado_counts.values[::-1], color=colors_est[::-1], edgecolor='black', linewidth=0.5)
    ax_d1.set_title("1. Reservas por Estado", fontsize=12, fontweight='bold', color='#1B365D')
    ax_d1.set_xlabel("N° Reservas", fontsize=9)
    for b in ax_d1.patches:
        w = b.get_width()
        ax_d1.text(w + 0.4, b.get_y() + b.get_height()/2, f"{int(w)}", va='center', fontsize=8, fontweight='bold')

    # Subplot 2: Medios de pago
    ax_d2 = plt.subplot2grid((2, 3), (0, 1))
    ax_d2.bar(pagos_por_medio.index, pagos_por_medio.values, color=colors_mp[:len(pagos_por_medio)], edgecolor='black', linewidth=0.5)
    ax_d2.set_title("2. Ingresos por Medio de Pago (S/.)", fontsize=12, fontweight='bold', color='#1B365D')
    ax_d2.set_xticks(range(len(pagos_por_medio.index)))
    ax_d2.set_xticklabels(pagos_por_medio.index, rotation=30, ha='right', fontsize=8)
    for b in ax_d2.patches:
        h = b.get_height()
        ax_d2.text(b.get_x() + b.get_width()/2, h + 2500, f"S/.{h/1000:.0f}k", ha='center', fontsize=8, fontweight='bold')

    # Subplot 3: Tendencia por mes
    ax_d3 = plt.subplot2grid((2, 3), (0, 2))
    ax_d3.plot(reservas_periodo.index, reservas_periodo.values, marker='o', color='#1B365D', linewidth=2)
    ax_d3.bar(reservas_periodo.index, reservas_periodo.values, color='#80DEEA', alpha=0.6, width=0.4)
    ax_d3.set_title("3. Reservas por Período / Fecha", fontsize=12, fontweight='bold', color='#1B365D')
    ax_d3.set_xlabel("Mes (2026)", fontsize=9)
    for idx, (p, val) in enumerate(reservas_periodo.items()):
        ax_d3.text(idx, val + 1, f"{val}", ha='center', fontsize=9, fontweight='bold')

    # Subplot 4: Top reservas clientes
    ax_d4 = plt.subplot2grid((2, 3), (1, 0))
    ax_d4.barh(top_reservas_cli['cliente_nombre'][::-1], top_reservas_cli['num_reservas'][::-1], color='#2E7D32', edgecolor='black', linewidth=0.5)
    ax_d4.set_title("4. Top 10 Clientes por Frecuencia", fontsize=12, fontweight='bold', color='#1B365D')
    ax_d4.set_xlabel("N° Reservas", fontsize=9)
    for b in ax_d4.patches:
        w = b.get_width()
        ax_d4.text(w + 0.1, b.get_y() + b.get_height()/2, f"{int(w)}", va='center', fontsize=8, fontweight='bold')

    # Subplot 5: Top ingresos clientes
    ax_d5 = plt.subplot2grid((2, 3), (1, 1))
    ax_d5.barh(ingresos_cli['cliente_nombre'][::-1], ingresos_cli['total_facturado'][::-1], color='#E65100', edgecolor='black', linewidth=0.5)
    ax_d5.set_title("5. Ingresos Totales por Cliente (S/.)", fontsize=12, fontweight='bold', color='#1B365D')
    ax_d5.set_xlabel("Soles (S/.)", fontsize=9)
    for b in ax_d5.patches:
        w = b.get_width()
        ax_d5.text(w + 200, b.get_y() + b.get_height()/2, f"S/.{w/1000:.1f}k", va='center', fontsize=8, fontweight='bold')

    # Subplot 6: Tarjeta Resumen de KPIs
    ax_d6 = plt.subplot2grid((2, 3), (1, 2))
    ax_d6.axis('off')
    kpi_card_text = (
        "╔═══════════════════════════════════════════╗\n"
        "║          INDICADORES CLAVE (KPIS)         ║\n"
        "╠═══════════════════════════════════════════╣\n"
        f"║  • TOTAL CLIENTES:        {len(df_clientes):>13,}   ║\n"
        f"║  • TOTAL RESERVAS:        {len(df_reservas):>13,}   ║\n"
        f"║  • TOTAL INGRESOS:   S/. {df_pagos['monto'].sum():>13,.2f} ║\n"
        f"║  • TICKET PROMEDIO:  S/. {df_pagos['monto'].sum()/len(df_reservas):>13,.2f} ║\n"
        "╠═══════════════════════════════════════════╣\n"
        f"║  • FACTURADO RESERVAS: S/. {df_reservas['precio_total'].sum():>10,.2f} ║\n"
        f"║  • PAGOS PROCESADOS:      {len(df_pagos):>13,}   ║\n"
        "╚═══════════════════════════════════════════╝\n"
    )
    ax_d6.text(0.05, 0.5, kpi_card_text, fontfamily='monospace', fontsize=11, 
               verticalalignment='center', color='#1B365D',
               bbox=dict(boxstyle='round,pad=1', facecolor='#E3F2FD', edgecolor='#1565C0', linewidth=2))

    plt.tight_layout(rect=[0, 0, 1, 0.94])
    reporte_png_path = os.path.join(EVIDENCIAS_DIR, "reporte.png")
    fig_all.savefig(reporte_png_path, facecolor=fig_all.get_facecolor(), edgecolor='none')
    plt.close(fig_all)
    print(f" -> Panel completo guardado exitosamente en: {reporte_png_path}")

    return {
        "chart1": chart1_path,
        "chart2": chart2_path,
        "chart3": chart3_path,
        "chart4": chart4_path,
        "chart5": chart5_path,
        "estado_counts": estado_counts,
        "pagos_por_medio": pagos_por_medio,
        "reservas_periodo": reservas_periodo,
        "top_reservas_cli": top_reservas_cli,
        "ingresos_cli": ingresos_cli
    }

# ------------------------------------------------------------------------------
# 6. COMPILACIÓN DE INFORME EJECUTIVO EN PDF (REPORTLAB)
# ------------------------------------------------------------------------------
def build_pdf_report(kpis, chart_data):
    pdf_filename = os.path.join(REPORTES_DIR, "reportes.pdf")
    print(f"[PDF] Compilando informe ejecutivo formal en: {pdf_filename}...")

    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Tipografía estricta: Times-Roman / Times-Bold, tamaño 12, color negro
    univ_style = ParagraphStyle(
        'CoverUniv',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        alignment=1,
        spaceAfter=15
    )

    title_style = ParagraphStyle(
        'CoverTitle',
        parent=styles['Heading1'],
        fontName='Times-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        alignment=1,
        spaceAfter=15
    )

    subtitle_style = ParagraphStyle(
        'CoverSubtitle',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        alignment=1,
        spaceAfter=30
    )

    h1_style = ParagraphStyle(
        'CustomH1',
        parent=styles['Heading1'],
        fontName='Times-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        spaceAfter=8
    )

    conclusion_style = ParagraphStyle(
        'ConclusionItem',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12,
        leading=16,
        textColor=colors.black,
        spaceAfter=10
    )

    story = []

    # ==========================================================================
    # PÁGINA 1: PORTADA SIMPLIFICADA
    # ==========================================================================
    story.append(Spacer(1, 100))
    story.append(Paragraph("UNIVERSIDAD NACIONAL DE CAJAMARCA", univ_style))
    story.append(Spacer(1, 15))
    story.append(Paragraph("INFORME EJECUTIVO ANALÍTICO DE INTELIGENCIA DE NEGOCIOS Y GESTIÓN DE SEGURIDAD", title_style))
    story.append(Paragraph("Tercera Evaluación de Base de Datos II - Grupo B (Alternativa B: Python)", subtitle_style))
    story.append(Spacer(1, 30))

    meta_table_data = [
        [Paragraph("<b>Asignatura:</b>", body_style), Paragraph("Base de Datos II", body_style)],
        [Paragraph("<b>Estudiante:</b>", body_style), Paragraph("Cristian Paul Alvarado Minchan / VonRabbid", body_style)],
        [Paragraph("<b>Base de Datos:</b>", body_style), Paragraph("TURISMOPERU_CPAM", body_style)],
    ]
    meta_table = Table(meta_table_data, colWidths=[150, 320])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.white),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('FONTNAME', (0, 0), (-1, -1), 'Times-Roman'),
        ('FONTSIZE', (0, 0), (-1, -1), 12),
        ('PADDING', (0, 0), (-1, -1), 8),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.black),
    ]))
    story.append(meta_table)
    story.append(PageBreak())

    # ==========================================================================
    # PÁGINA 2: RESUMEN EJECUTIVO Y TABLA DE KPIS
    # ==========================================================================
    story.append(Paragraph("1. Resumen Ejecutivo de Métricas Clave (KPIs)", h1_style))
    story.append(Paragraph(
        "El presente informe recopila los hallazgos analíticos generados a partir de la explotación de la base de datos transaccional <b>TURISMOPERU_CPAM</b>. Los indicadores reflejan la consolidación del primer semestre operativo del año 2026, abarcando clientes registrados, reservas concretadas y la totalidad de los flujos de pagos liquidados.",
        body_style
    ))
    story.append(Spacer(1, 8))

    kpi_table_data = [
        ["Indicador de Rendimiento (KPI)", "Valor Obtenido", "Unidad de Medida", "Interpretación de Negocio"],
        ["Total Clientes", f"{kpis['Total Clientes']:,}", "Clientes", "Cartera de clientes registrados y activos."],
        ["Total Reservas", f"{kpis['Total Reservas']:,}", "Transacciones", "Volumen total de expedientes de reserva creados."],
        ["Total Ingresos Cobrados", f"S/. {kpis['Total Ingresos']:,.2f}", "Soles (PEN)", "Monto neto liquidado a través de pasarelas y cajas."],
        ["Ticket Promedio (Cobrado)", f"S/. {kpis['Ticket Promedio']:,.2f}", "Soles / Reserva", "Recaudación media efectiva por reserva generada."],
        ["Facturación Contratada", f"S/. {kpis['Facturación Contratada']:,.2f}", "Soles (PEN)", "Valor total comercial de paquetes y reservas emitidas."]
    ]
    t_kpi = Table(kpi_table_data, colWidths=[130, 95, 90, 189])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Times-Bold'),
        ('FONTNAME', (0, 1), (-1, -1), 'Times-Roman'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('ALIGN', (1, 1), (1, -1), 'RIGHT'),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 15))

    # Gráfico 1: Reservas por Estado (a todo color)
    story.append(Paragraph("2. Análisis Operativo: Reservas por Estado", h1_style))
    story.append(Paragraph(
        "Se evalúa el ciclo de vida de los contratos de reserva turística para determinar la efectividad en el cumplimiento y detección temprana de deserciones.",
        body_style
    ))
    story.append(RLImage(chart_data['chart1'], width=6.5*inch, height=3.2*inch))
    story.append(PageBreak())

    # ==========================================================================
    # PÁGINA 3: INGRESOS POR MEDIO DE PAGO
    # ==========================================================================
    story.append(Paragraph("3. Análisis Financiero: Ingresos por Medio de Pago", h1_style))
    story.append(Paragraph(
        "La recaudación acumulada de <b>S/. 351,975.00</b> fue procesada mediante múltiples canales financieros. A continuación se detalla la concentración del capital según pasarela electrónica, billetera móvil y ventanilla:",
        body_style
    ))
    story.append(Spacer(1, 5))

    p_df = chart_data['pagos_por_medio']
    mp_table_data = [["Medio de Pago", "Monto Recaudado (S/.)", "Participación (%)", "Transacciones"]]
    for mp_name, monto in p_df.items():
        pct = (monto / p_df.sum()) * 100
        mp_table_data.append([mp_name, f"S/. {monto:,.2f}", f"{pct:.2f}%", "Procesadas"])
    
    t_mp = Table(mp_table_data, colWidths=[150, 130, 110, 114])
    t_mp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Times-Bold'),
        ('FONTNAME', (0, 1), (-1, -1), 'Times-Roman'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('ALIGN', (1, 1), (2, -1), 'RIGHT'),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_mp)
    story.append(Spacer(1, 10))

    # Gráfico 2: Ingresos por Medio de Pago (a todo color)
    story.append(RLImage(chart_data['chart2'], width=6.5*inch, height=3.3*inch))
    story.append(PageBreak())

    # ==========================================================================
    # PÁGINA 4: ESTACIONALIDAD TEMPORAL Y COMPORTAMIENTO DE CLIENTES
    # ==========================================================================
    story.append(Paragraph("4. Tendencia Temporal y Estacionalidad de la Demanda", h1_style))
    story.append(Paragraph(
        "El análisis mensual evidencia una clara aceleración en la adquisición de paquetes turísticos conforme avanza el calendario hacia la temporada alta de invierno andino y fiestas nacionales.",
        body_style
    ))
    story.append(RLImage(chart_data['chart3'], width=6.5*inch, height=2.8*inch))
    story.append(Spacer(1, 10))

    story.append(Paragraph("5. Segmentación de Clientes: Frecuencia y Concentración de Valor", h1_style))
    story.append(Paragraph(
        "A continuación se contrastan los clientes con mayor frecuencia de contratación frente a aquellos con mayor volumen de facturación monetaria:",
        body_style
    ))
    story.append(RLImage(chart_data['chart4'], width=6.5*inch, height=2.6*inch))
    story.append(PageBreak())

    story.append(Paragraph("6. Facturación Acumulada por Cliente", h1_style))
    story.append(RLImage(chart_data['chart5'], width=6.5*inch, height=3.3*inch))
    story.append(Spacer(1, 15))

    # ==========================================================================
    # PÁGINA 5 / 6: CONCLUSIONES ANALÍTICAS FUNDAMENTADAS
    # ==========================================================================
    story.append(Paragraph("7. Conclusiones Analíticas Fundamentadas", h1_style))
    story.append(Paragraph(
        "A partir de las métricas obtenidas y de la explotación estadística de los registros de <b>TURISMOPERU_CPAM</b>, se formulan las siguientes 5 conclusiones analíticas institucionales:",
        body_style
    ))
    story.append(Spacer(1, 8))

    c1 = (
        "<b>1. Se observa que</b> el <b>36%</b> de las reservas emitidas se encuentran en estado <b>'Completada'</b> "
        "(36 de 100 reservas), en tanto que un <b>20%</b> adicional permanece activo en estado 'Parcialmente Pagada' o 'En Proceso'. "
        "Asimismo, se constata una tasa de cancelación sumamente reducida de solo el <b>6%</b> (6 casos), "
        "lo que ratifica una alta eficiencia en la conversión comercial y estabilidad en el cumplimiento de los contratos turísticos."
    )
    story.append(Paragraph(c1, conclusion_style))

    c2 = (
        "<b>2. El medio de pago con mayor</b> recaudación absoluta es la pasarela de tarjetas <b>Visa</b> con un total de "
        "<b>S/. 148,500.00</b> (representando el <b>42.19%</b> de los fondos totales), seguido por <b>Mastercard</b> con "
        "<b>S/. 84,025.00</b> (23.87%) y las <b>Transferencias Bancarias</b> con <b>S/. 51,725.00</b> (14.70%). "
        "En conjunto, las transacciones electrónicas bancarizadas concentran más del 80% de las operaciones, "
        "desplazando ampliamente al uso de dinero en efectivo (S/. 18,450.00 / 5.24%)."
    )
    story.append(Paragraph(c2, conclusion_style))

    c3 = (
        "<b>3. El periodo con mayor</b> afluencia de reservas corresponde a los meses de <b>mayo (25 reservas)</b> y "
        "<b>junio de 2026 (26 reservas)</b>, los cuales concentran en conjunto el <b>51%</b> de todo el movimiento "
        "registrado en el semestre. Este patrón demuestra la marcada estacionalidad de la demanda hacia festividades de invierno, "
        "vacaciones de medio año e itinerarios cusqueños y de sierra peruana, en contraste con el inicio de año (3 reservas en enero)."
    )
    story.append(Paragraph(c3, conclusion_style))

    c4 = (
        "<b>4. Los clientes que concentran</b> el mayor nivel de facturación monetaria son <b>Luisa María Herrera</b> "
        "(S/. 19,200.00) y <b>Yolanda Marín</b> (S/. 18,800.00), mientras que el cliente con mayor frecuencia de contratación "
        "es <b>Raúl Figueroa</b> con un acumulado de 4 reservas. Esto demuestra la presencia de un segmento VIP corporativo / familiar "
        "que representa un alto valor de vida del cliente (Customer Lifetime Value - CLV)."
    )
    story.append(Paragraph(c4, conclusion_style))

    c5 = (
        "<b>5. El comportamiento de</b> ticket promedio, situado en <b>S/. 3,519.75 cobrado por reserva</b> "
        "(y S/. 4,071.00 en valor contractual bruto contratado), evidencia que los usuarios contratan paquetes con "
        "servicios combinados de hospedaje premium y excursiones turísticas de estancia prolongada, "
        "lo que brinda una sólida base financiera y valida la viabilidad comercial de la plataforma TurismoPerú."
    )
    story.append(Paragraph(c5, conclusion_style))

    # Construir documento sin NumberedCanvas
    doc.build(story)
    print(f"[PDF] Informe formal generado exitosamente en: {pdf_filename}")

# ------------------------------------------------------------------------------
# 7. COMPILACIÓN DE INFORME EJECUTIVO EN WORD (.DOCX)
# ------------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_cell_text(cell, text, bold=False, size_pt=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    return run

def add_custom_heading(doc, text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True
    return h

def add_custom_paragraph(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
        r_pre.bold = True
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_docx_report(kpis, chart_data):
    docx_filename = os.path.join(REPORTES_DIR, "reportes.docx")
    print(f"[DOCX] Compilando informe ejecutivo formal en: {docx_filename}...")

    doc = Document()
    
    # Configurar estilo 'Normal' estrictamente a Times New Roman 12 pt, color negro
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # Configurar márgenes de página
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ==========================================================================
    # PORTADA SIMPLIFICADA (DOCX)
    # ==========================================================================
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(80)

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_after = Pt(15)
    r_univ = p_univ.add_run("UNIVERSIDAD NACIONAL DE CAJAMARCA\n")
    r_univ.bold = True
    r_univ.font.name = "Times New Roman"
    r_univ.font.size = Pt(12)
    r_univ.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(15)
    r_title = p_title.add_run("INFORME EJECUTIVO ANALÍTICO DE INTELIGENCIA DE NEGOCIOS Y GESTIÓN DE SEGURIDAD\n")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(12)
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(30)
    r_sub = p_sub.add_run("Tercera Evaluación de Base de Datos II - Grupo B (Alternativa B: Python)\n")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(0, 0, 0)

    # Tabla de portada: ÚNICAMENTE 3 FILAS (sin colores de fondo, con texto en Times New Roman 12 y bordes negros simples)
    table_meta = doc.add_table(rows=3, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_meta, color="000000", sz="4", val="single")

    meta_rows = [
        ("Asignatura:", "Base de Datos II"),
        ("Estudiante:", "Cristian Paul Alvarado Minchan / VonRabbid"),
        ("Base de Datos:", "TURISMOPERU_CPAM"),
    ]
    for idx, (k, v) in enumerate(meta_rows):
        row = table_meta.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        format_cell_text(cell_k, k, bold=True, size_pt=12)
        format_cell_text(cell_v, v, bold=False, size_pt=12)
        cell_k.width = Inches(2.0)
        cell_v.width = Inches(4.5)

    doc.add_page_break()

    # ==========================================================================
    # 1. RESUMEN EJECUTIVO Y TABLA DE KPIS
    # ==========================================================================
    add_custom_heading(doc, "1. Resumen Ejecutivo de Métricas Clave (KPIs)", level=1)
    add_custom_paragraph(
        doc,
        "El presente informe recopila los hallazgos analíticos generados a partir de la explotación de la base de datos transaccional TURISMOPERU_CPAM. "
        "Los indicadores reflejan la consolidación del primer semestre operativo del año 2026, abarcando clientes registrados, reservas concretadas y flujos de pagos liquidados."
    )

    t_kpis = doc.add_table(rows=6, cols=4)
    t_kpis.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_kpis, color="000000", sz="4", val="single")

    kpi_headers = ["Indicador Clave (KPI)", "Valor Obtenido", "Unidad", "Interpretación de Negocio"]
    for c_idx, h_text in enumerate(kpi_headers):
        cell = t_kpis.rows[0].cells[c_idx]
        format_cell_text(cell, h_text, bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_background(cell, "F2F2F2")

    kpi_rows_data = [
        ("Total Clientes", f"{kpis['Total Clientes']:,}", "Clientes", "Cartera total de clientes registrados y activos."),
        ("Total Reservas", f"{kpis['Total Reservas']:,}", "Contratos", "Volumen total de expedientes de reserva creados."),
        ("Total Ingresos Cobrados", f"S/. {kpis['Total Ingresos']:,.2f}", "Soles (PEN)", "Monto neto liquidado a través de pasarelas y cajas."),
        ("Ticket Promedio (Cobrado)", f"S/. {kpis['Ticket Promedio']:,.2f}", "Soles / Reserva", "Recaudación media efectiva por reserva generada."),
        ("Facturación Contratada", f"S/. {kpis['Facturación Contratada']:,.2f}", "Soles (PEN)", "Valor total comercial de paquetes y reservas emitidas.")
    ]
    for r_idx, (k, val, un, desc) in enumerate(kpi_rows_data):
        row = t_kpis.rows[r_idx + 1]
        format_cell_text(row.cells[0], k, bold=False, size_pt=11)
        format_cell_text(row.cells[1], val, bold=False, size_pt=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
        format_cell_text(row.cells[2], un, bold=False, size_pt=11)
        format_cell_text(row.cells[3], desc, bold=False, size_pt=11)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)

    # 2. Reservas por Estado
    add_custom_heading(doc, "2. Análisis Operativo: Reservas por Estado", level=1)
    add_custom_paragraph(
        doc,
        "Se evalúa la distribución del ciclo de vida de los contratos de reserva turística para determinar la efectividad en el cumplimiento y detección de deserciones."
    )
    if os.path.exists(chart_data['chart1']):
        doc.add_picture(chart_data['chart1'], width=Inches(6.2))

    doc.add_page_break()

    # ==========================================================================
    # 3. INGRESOS POR MEDIO DE PAGO
    # ==========================================================================
    add_custom_heading(doc, "3. Análisis Financiero: Ingresos por Medio de Pago", level=1)
    add_custom_paragraph(
        doc,
        "La recaudación acumulada fue procesada a través de diversos canales financieros y pasarelas de pago digitales. "
        "A continuación se presenta el desglose del capital según canal utilizado:"
    )
    
    p_df = chart_data['pagos_por_medio']
    t_mp = doc.add_table(rows=len(p_df) + 1, cols=3)
    t_mp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mp, color="000000", sz="4", val="single")

    mp_headers = ["Medio de Pago", "Monto Recaudado (S/.)", "Participación (%)"]
    for c_idx, h_text in enumerate(mp_headers):
        cell = t_mp.rows[0].cells[c_idx]
        format_cell_text(cell, h_text, bold=True, size_pt=11, align=WD_ALIGN_PARAGRAPH.RIGHT if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_background(cell, "F2F2F2")

    for r_idx, (mp_name, monto) in enumerate(p_df.items()):
        pct = (monto / p_df.sum()) * 100
        row = t_mp.rows[r_idx + 1]
        format_cell_text(row.cells[0], mp_name, bold=False, size_pt=11)
        format_cell_text(row.cells[1], f"S/. {monto:,.2f}", bold=False, size_pt=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
        format_cell_text(row.cells[2], f"{pct:.2f}%", bold=False, size_pt=11, align=WD_ALIGN_PARAGRAPH.RIGHT)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(8)
    if os.path.exists(chart_data['chart2']):
        doc.add_picture(chart_data['chart2'], width=Inches(6.2))

    doc.add_page_break()

    # ==========================================================================
    # 4. TENDENCIA TEMPORAL Y SEGMENTACIÓN DE CLIENTES
    # ==========================================================================
    add_custom_heading(doc, "4. Tendencia Temporal y Estacionalidad de la Demanda", level=1)
    add_custom_paragraph(
        doc,
        "El análisis mensual evidencia una clara aceleración en la adquisición de paquetes turísticos conforme avanza el calendario hacia la temporada alta de invierno andino y fiestas nacionales."
    )
    if os.path.exists(chart_data['chart3']):
        doc.add_picture(chart_data['chart3'], width=Inches(6.2))

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(8)

    add_custom_heading(doc, "5. Segmentación de Clientes: Frecuencia y Concentración de Valor", level=1)
    add_custom_paragraph(
        doc,
        "A continuación se contrastan los clientes con mayor frecuencia de contratación frente a aquellos con mayor volumen de facturación monetaria:"
    )
    if os.path.exists(chart_data['chart4']):
        doc.add_picture(chart_data['chart4'], width=Inches(6.2))
    
    doc.add_paragraph()
    if os.path.exists(chart_data['chart5']):
        doc.add_picture(chart_data['chart5'], width=Inches(6.2))

    doc.add_page_break()

    # ==========================================================================
    # 6. CONCLUSIONES ANALÍTICAS FUNDAMENTADAS
    # ==========================================================================
    add_custom_heading(doc, "6. Conclusiones Analíticas Fundamentadas", level=1)
    add_custom_paragraph(
        doc,
        "A partir de las métricas obtenidas y de la explotación estadística de los registros de TURISMOPERU_CPAM, se formulan las siguientes 5 conclusiones analíticas institucionales:"
    )

    conclusiones = [
        "1. Se observa que el 36% de las reservas emitidas se encuentran en estado 'Completada' (36 de 100 reservas), en tanto que un 20% adicional permanece activo en estado 'Parcialmente Pagada' o 'En Proceso'. Asimismo, se constata una tasa de cancelación sumamente reducida de solo el 6% (6 casos), lo que ratifica una alta eficiencia en la conversión comercial y estabilidad en el cumplimiento de los contratos turísticos.",
        "2. El medio de pago con mayor recaudación absoluta es la pasarela de tarjetas Visa con un total de S/. 148,500.00 (representando el 42.19% de los fondos totales), seguido por Mastercard con S/. 84,025.00 (23.87%) y las Transferencias Bancarias con S/. 51,725.00 (14.70%). En conjunto, las transacciones electrónicas bancarizadas concentran más del 80% de las operaciones, desplazando ampliamente al uso de dinero en efectivo (S/. 18,450.00 / 5.24%).",
        "3. El periodo con mayor afluencia de reservas corresponde a los meses de mayo (25 reservas) y junio de 2026 (26 reservas), los cuales concentran en conjunto el 51% de todo el movimiento registrado en el semestre. Este patrón demuestra la marcada estacionalidad de la demanda hacia festividades de invierno, vacaciones de medio año e itinerarios cusqueños y de sierra peruana, en contraste con el inicio de año (3 reservas en enero).",
        "4. Los clientes que concentran el mayor nivel de facturación monetaria son Luisa María Herrera (S/. 19,200.00) y Yolanda Marín (S/. 18,800.00), mientras que el cliente con mayor frecuencia de contratación es Raúl Figueroa con un acumulado de 4 reservas. Esto demuestra la presencia de un segmento VIP corporativo / familiar que representa un alto valor de vida del cliente (Customer Lifetime Value - CLV).",
        "5. El comportamiento de ticket promedio, situado en S/. 3,519.75 cobrado por reserva (y S/. 4,071.00 en valor contractual bruto contratado), evidencia que los usuarios contratan paquetes con servicios combinados de hospedaje premium y excursiones turísticas de estancia prolongada, lo que brinda una sólida base financiera y valida la viabilidad comercial de la plataforma TurismoPerú."
    ]

    for c in conclusiones:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_after = Pt(8)
        p_c.paragraph_format.line_spacing = 1.15
        run_c = p_c.add_run(c)
        run_c.font.name = "Times New Roman"
        run_c.font.size = Pt(12)
        run_c.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(docx_filename)
    print(f"[DOCX] Informe Word generado exitosamente en: {docx_filename}")

# ------------------------------------------------------------------------------
# 8. FUNCIÓN PRINCIPAL DE EJECUCIÓN
# ------------------------------------------------------------------------------
def main():
    print("================================================================================")
    print("  TURISMOPERÚ - EJECUCIÓN DE MÓDULO ANALÍTICO EN PYTHON (ALTERNATIVA B)")
    print("  Universidad Nacional de Cajamarca | Cristian Paul Alvarado Minchan / VonRabbid")
    print("================================================================================")
    
    try:
        engine = get_db_engine()
        df_clientes, df_reservas, df_pagos = extract_data(engine)
        kpis = compute_kpis(df_clientes, df_reservas, df_pagos)
        chart_data = generate_charts(df_clientes, df_reservas, df_pagos)
        build_pdf_report(kpis, chart_data)
        build_docx_report(kpis, chart_data)
        
        # Limpieza de imágenes temporales
        import shutil
        if os.path.exists(TEMP_IMG_DIR):
            shutil.rmtree(TEMP_IMG_DIR, ignore_errors=True)
            
        print("\n[ÉXITO] Todo el pipeline analítico y documental (PDF y DOCX) fue ejecutado satisfactoriamente.")
    except Exception as e:
        print(f"\n[ERROR CRÍTICO] Falló la ejecución: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == "__main__":
    main()
